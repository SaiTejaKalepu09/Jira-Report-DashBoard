from io import BytesIO

from docx import Document
from docx.shared import Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT


GREEN = "DAF2D0"
HEADER_BLUE = "#1E65B8"
ROW_LIGHT_BLUE = "#579CE1"

def format_summary_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):

            # Header row OR first column
            if row_idx == 0 or col_idx == 0:
                set_cell_color(
                    cell,
                    HEADER_BLUE
                )
            else:
                set_cell_color(
                    cell,
                    ROW_LIGHT_BLUE
                )

            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(
                        0,
                        0,
                        0
                    )

                    if row_idx == 0:
                        run.bold = True

def set_cell_color(cell, color):

    tc_pr = cell._tc.get_or_add_tcPr()

    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color}"/>'
    )

    tc_pr.append(shading)


def format_table(table):

    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for row in table.rows:

        for cell in row.cells:

            set_cell_color(
                cell,
                GREEN
            )

            for paragraph in cell.paragraphs:

                for run in paragraph.runs:

                    run.font.color.rgb = RGBColor(
                        0,
                        0,
                        0
                    )


def create_word_report(
    chart_png,
    requirements_df,
    other_df
):

    document = Document()

    # ==========================
    # Summary Table
    # ==========================

    summary_data = [
        [
            "Configuration",
            ""
        ],
        [
            "Integration",
            ""
        ],
        [
            "Portal",
            "37 USC in Sprint 9, including 28 in Testing. Defects : Total 20, 5 planned today, 7 planned tomorrow, 8 in analysis."
        ],
        [
            "Conversion",
            ""
        ],
        [
            "Risks/Dependencies",
            ""
        ]
    ]

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.style = "Table Grid"

    header = table.rows[0].cells

    header[0].text = "Area"
    header[1].text = "Update"

    for item in summary_data:

        row = table.add_row().cells

        row[0].text = item[0]
        row[1].text = item[1]

    format_summary_table(table)

    document.add_paragraph()

    # ==========================
    # Burn Down Chart
    # ==========================

    document.add_heading(
        "Portal Team",
        level=2
    )

    img_stream = BytesIO(
        chart_png
    )

    document.add_picture(
        img_stream,
        width=Inches(6.5)
    )

    # ==========================
    # Other Statuses
    # ==========================

    document.add_heading(
        "List of Items due today & next few days: ",
        level=2
    )

    other_table = document.add_table(
        rows=1,
        cols=len(other_df.columns)
    )

    other_table.style = "Table Grid"

    for idx, col in enumerate(
        other_df.columns
    ):
        other_table.rows[0].cells[idx].text = str(col)

    for _, row_data in other_df.iterrows():

        row = other_table.add_row().cells

        for idx, value in enumerate(
            row_data
        ):
            row[idx].text = str(value)

    format_table(
        other_table
    )

    document.add_paragraph()

    # ==========================
    # Requirements
    # ==========================

    document.add_heading(
        "Defects in Requirements status",
        level=2
    )

    req_table = document.add_table(
        rows=1,
        cols=len(
            requirements_df.columns
        )
    )

    req_table.style = "Table Grid"

    for idx, col in enumerate(
        requirements_df.columns
    ):
        req_table.rows[0].cells[idx].text = str(col)

    for _, row_data in requirements_df.iterrows():

        row = req_table.add_row().cells

        for idx, value in enumerate(
            row_data
        ):
            row[idx].text = str(value)

    format_table(
        req_table
    )

    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output.getvalue()