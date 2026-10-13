from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import (
    PatternFill,
    Font,
    Border,
    Side,
    Alignment
)


def style_sheet(ws):

    header_fill = PatternFill(
        fill_type="solid",
        start_color="DAF2D0",
        end_color="DAF2D0"
    )

    first_column_fill = PatternFill(
        fill_type="solid",
        start_color="DAF2D0",
        end_color="DAF2D0"
    )

    white_fill = PatternFill(
        fill_type="solid",
        start_color="DAF2D0",
        end_color="DAF2D0"
    )

    gray_fill = PatternFill(
        fill_type="solid",
        start_color="DAF2D0",
        end_color="DAF2D0"
    )

    header_font = Font(
        color="000000",
        bold=True
    )

    normal_font = Font(
        color="000000"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for row_num in range(2, ws.max_row + 1):

        row_fill = (
            white_fill
            if row_num % 2 == 0
            else gray_fill
        )

        for col_num in range(1, ws.max_column + 1):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            if col_num == 1:
                cell.fill = first_column_fill
            else:
                cell.fill = row_fill

            cell.font = normal_font

    for column_cells in ws.columns:

        max_length = 0

        column_letter = (
            column_cells[0].column_letter
        )

        for cell in column_cells:

            try:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )
            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = max_length + 5


def create_full_report(
    chart_png,
    requirements_df,
    other_df
):

    wb = Workbook()

    # ==================================================
    # Sheet 1 - Burn Down Chart
    # ==================================================

    ws_chart = wb.active
    ws_chart.title = "Burn_Down_Chart"

    image_stream = BytesIO(chart_png)

    img = Image(image_stream)

    img.width = 900
    img.height = 450

    ws_chart.add_image(
        img,
        "A1"
    )

    # ==================================================
    # Sheet 2 - Other Statuses
    # ==================================================

    ws_other = wb.create_sheet(
        "Other_Statuses"
    )

    ws_other.append(
        other_df.columns.tolist()
    )

    for row in other_df.values.tolist():
        ws_other.append(row)

    style_sheet(ws_other)

    # ==================================================
    # Sheet 3 - Requirements
    # ==================================================

    ws_req = wb.create_sheet(
        "Requirements"
    )

    ws_req.append(
        requirements_df.columns.tolist()
    )

    for row in requirements_df.values.tolist():
        ws_req.append(row)

    style_sheet(ws_req)

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()